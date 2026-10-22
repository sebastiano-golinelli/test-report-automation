"""Excel report: summary per material, all results, issues to review, chart."""

from statistics import mean

from openpyxl import Workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .specs import SPECS

NAVY, GREY = "1F2A44", "F2F4F8"
FILL = {"PASS": "D9F2E3", "FAIL": "FAD4D4", "REVIEW": "FFF0C2"}
THIN = Side(style="thin", color="C8CFDB")


def _header(ws, row, values):
    for col, v in enumerate(values, 1):
        c = ws.cell(row=row, column=col, value=v)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _table(ws, first_row, rows):
    for i, values in enumerate(rows):
        for col, v in enumerate(values, 1):
            c = ws.cell(row=first_row + i, column=col, value=v)
            c.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
            c.alignment = Alignment(vertical="top", wrap_text=isinstance(v, str) and len(v) > 40)
            if i % 2:
                c.fill = PatternFill("solid", fgColor=GREY)


def _widths(ws, widths):
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w


def write_report(records, status, issues, path):
    wb = Workbook()

    # --- Summary -------------------------------------------------------------------------------
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Tensile test reports - automated summary"
    ws["A1"].font = Font(bold=True, size=16, color=NAVY)
    counts = {k: sum(1 for s, _ in status.values() if s == k) for k in ("PASS", "FAIL", "REVIEW")}
    ws["A2"] = (f"{len(records)} reports read: {counts['PASS']} pass, {counts['FAIL']} below specification, "
                f"{counts['REVIEW']} to review by hand (see the 'Needs review' sheet).")
    _header(ws, 4, ["Material", "Specimens passed", "Yield mean [MPa]", "Yield min [MPa]", "Spec min yield [MPa]",
                    "UTS mean [MPa]", "Spec min UTS [MPa]", "Elongation mean [%]", "Specification"])
    rows = []
    for mat, spec in SPECS.items():
        ok = [r for r in records if r.material == mat and status[r.file][0] in ("PASS", "FAIL")]
        if not ok:
            continue
        rows.append([mat, f"{sum(status[r.file][0] == 'PASS' for r in ok)} / {len(ok)}",
                     round(mean(r.yield_mpa for r in ok)), round(min(r.yield_mpa for r in ok)), spec["yield_min"],
                     round(mean(r.uts_mpa for r in ok)), spec["uts_min"],
                     round(mean(r.elongation_pct for r in ok), 1), spec["source"]])
    _table(ws, 5, rows)
    for i in range(len(rows)):
        ws.cell(row=5 + i, column=8).number_format = "0.0"
    ws.cell(row=6 + len(rows), column=1,
            value="Statistics use only checked results (PASS or FAIL); reports under review are excluded.").font = \
        Font(italic=True, color="666666")
    _widths(ws, [14, 16, 15, 14, 18, 14, 17, 18, 26])
    ws.row_dimensions[4].height = 32

    # --- Results -------------------------------------------------------------------------------
    wr = wb.create_sheet("Results")
    cols = ["Specimen", "Material", "Yield [MPa]", "UTS [MPa]", "Elongation [%]", "Test date", "Status", "Reason",
            "Read by", "Lab format", "Source file", "Remarks"]
    _header(wr, 1, cols)
    order = {"FAIL": 0, "REVIEW": 1, "PASS": 2}
    ordered = sorted(records, key=lambda r: (order[status[r.file][0]], r.specimen_id or ""))
    _table(wr, 2, [[r.specimen_id, r.material, r.yield_mpa, r.uts_mpa, r.elongation_pct, r.test_date,
                    status[r.file][0], status[r.file][1], r.method, r.layout, r.file, r.remarks] for r in ordered])
    for i, r in enumerate(ordered, 2):
        c = wr.cell(row=i, column=7)
        c.fill, c.font = PatternFill("solid", fgColor=FILL[status[r.file][0]]), Font(bold=True)
        for col, fmt in ((3, "0"), (4, "0"), (5, "0.0")):
            wr.cell(row=i, column=col).number_format = fmt
    wr.freeze_panes = "B2"
    wr.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(ordered) + 1}"
    _widths(wr, [11, 12, 11, 11, 13, 12, 10, 52, 18, 24, 30, 46])

    # --- Needs review --------------------------------------------------------------------------
    wn = wb.create_sheet("Needs review")
    _header(wn, 1, ["Specimen", "Source file", "Issue"])
    _table(wn, 2, [[i["specimen"], i["file"], i["issue"]] for i in issues])
    _widths(wn, [11, 30, 90])

    # --- AI evidence (only when the AI step ran) -------------------------------------------------
    ai_rows = [[r.specimen_id, r.file, f, q] for r in records for f, q in r.evidence.items()]
    if ai_rows:
        we = wb.create_sheet("AI evidence")
        _header(we, 1, ["Specimen", "Source file", "Field", "Quote from the report"])
        _table(we, 2, ai_rows)
        _widths(we, [11, 26, 16, 90])

    # --- Chart: yield vs UTS per material --------------------------------------------------------
    wc = wb.create_sheet("Chart")
    chart = ScatterChart()
    chart.title = "Yield vs tensile strength (checked results)"
    chart.x_axis.title, chart.y_axis.title = "Tensile strength [MPa]", "Yield strength [MPa]"
    chart.x_axis.delete = chart.y_axis.delete = False          # recent Excel hides the axes otherwise
    chart.x_axis.scaling.min, chart.y_axis.scaling.min = 200, 200
    chart.x_axis.majorGridlines = None
    chart.legend.position = "r"
    for title in (chart.title, chart.x_axis.title, chart.y_axis.title):
        title.overlay = False                                 # keep titles outside the plot area
    colours = iter(["2F6FD6", "E8743B", "8A5CD1"])
    chart.height, chart.width = 11, 18
    row = 1
    for mat in SPECS:
        pts = [r for r in records if r.material == mat and status[r.file][0] in ("PASS", "FAIL")]
        if not pts:
            continue
        wc.cell(row=row, column=1, value=mat)
        for k, r in enumerate(pts, 1):
            wc.cell(row=row + k, column=1, value=r.uts_mpa)
            wc.cell(row=row + k, column=2, value=r.yield_mpa)
        s = Series(Reference(wc, min_col=2, min_row=row + 1, max_row=row + len(pts)),
                   Reference(wc, min_col=1, min_row=row + 1, max_row=row + len(pts)), title=mat)
        s.marker.symbol, s.marker.size, s.graphicalProperties.line.noFill = "circle", 9, True
        colour = next(colours)
        s.marker.graphicalProperties.solidFill = colour
        s.marker.graphicalProperties.line.solidFill = colour
        chart.series.append(s)
        row += len(pts) + 2
    wc.add_chart(chart, "D2")

    wb.save(path)
    return path
