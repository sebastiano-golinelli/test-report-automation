"""The pipeline against the ground truth of the generated samples.

Run from the repository root:  pytest
"""

import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from reportbot.__main__ import run
from reportbot.ai import Extraction, Value, ai_complete
from reportbot.extract import extract

ROOT = Path(__file__).resolve().parents[1]
PDFS = ROOT / "samples" / "pdfs"
TRUTH = {t["file"]: t for t in json.loads((ROOT / "samples" / "ground_truth.json").read_text(encoding="utf-8"))}


@pytest.mark.parametrize("name", [f for f in TRUTH if not f.startswith("gamma")])
def test_rules_read_structured_reports_exactly(name):
    rec, t = extract(PDFS / name), TRUTH[name]
    assert (rec.specimen_id, rec.material, rec.test_date) == (t["id"], t["material"], t["test_date"])
    assert rec.yield_mpa == pytest.approx(t["yield_mpa"], abs=0.05)
    assert rec.uts_mpa == pytest.approx(t["uts_mpa"], abs=0.05)
    assert rec.elongation_pct == t["elongation_pct"]          # None where the lab left it blank


@pytest.mark.parametrize("name", [f for f in TRUTH if f.startswith("gamma")])
def test_free_text_is_not_guessed(name):
    rec = extract(PDFS / name)
    assert rec.specimen_id == TRUTH[name]["id"]
    assert rec.yield_mpa is None and rec.uts_mpa is None and rec.elongation_pct is None


def test_every_planted_problem_is_caught(tmp_path):
    records, status, issues = run(PDFS, tmp_path / "report.xlsx")
    by_file = {f: s for f, (s, _) in status.items()}
    assert by_file["alpha_AL-003.pdf"] == "FAIL"                       # elongation below spec
    assert by_file["beta_ST-103.pdf"] == "REVIEW"                      # elongation missing
    assert by_file["beta_TI-202.pdf"] == "REVIEW"                      # yield > UTS
    assert by_file["alpha_ST-101.pdf"] == by_file["alpha_ST-101_resubmitted.pdf"] == "REVIEW"   # duplicate
    assert all(by_file[f] == "REVIEW" for f in by_file if f.startswith("gamma"))   # free text, no AI
    passed = {f for f, s in by_file.items() if s == "PASS"}
    assert passed == {"alpha_AL-001.pdf", "alpha_AL-002.pdf", "alpha_ST-102.pdf", "beta_AL-004.pdf",
                      "beta_TI-201.pdf"}
    wb = load_workbook(tmp_path / "report.xlsx")
    assert wb.sheetnames == ["Summary", "Results", "Needs review", "Chart"]
    assert wb["Needs review"].max_row - 1 == len(issues)


class FakeClient:
    """Stands in for anthropic.Anthropic(): returns a prepared Extraction."""

    def __init__(self, extraction):
        parse = lambda **kwargs: SimpleNamespace(stop_reason="end_turn", parsed_output=extraction)  # noqa: E731
        self.beta = SimpleNamespace(messages=SimpleNamespace(parse=parse))


def test_ai_values_are_kept_only_with_evidence_in_the_report():
    rec = extract(PDFS / "gamma_TI-203.pdf")
    answer = Extraction(
        specimen_id="TI-203", material="Ti-6Al-4V", test_date="2026-10-03",
        yield_strength_mpa=Value(value=905, quote="Plastic deformation began at roughly 905 MPa"),
        tensile_strength_mpa=Value(value=968, quote="failed at a peak stress of 968 MPa"),
        elongation_pct=Value(value=15.0, quote="after stretching by 15 percent"))       # invented: not in text
    notes = ai_complete(rec, FakeClient(answer))
    assert (rec.yield_mpa, rec.uts_mpa) == (905, 968)
    assert rec.elongation_pct is None and "elongation_pct" not in rec.evidence
    assert any("not found in the report" in n for n in notes)
    assert rec.missing == ["elongation_pct"]
